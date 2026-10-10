import json
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

from moneyhub.pipeline import Request, generate
from moneyhub.store import Store


@pytest.fixture
def generated():
    msgs = []
    res = generate(Request("teacher", "card_deck", sizes=["letter"], providers=["builtin"], seed=3), progress=msgs.append)
    return res, msgs


def test_pipeline_produces_complete_upload_package(generated):
    res, msgs = generated
    f = res.folder
    assert (f / "product").glob("*.pdf") and list((f / "product").glob("*.pdf"))
    etsy = sorted(p.name for p in (f / "images" / "etsy").glob("*.jpg"))
    assert len(etsy) == 9 and sum("mockup" in n for n in etsy) == 3   # six listing images plus three "in use" mockups
    assert len(list((f / "images" / "pinterest").glob("*.png"))) == 6
    for m in ("etsy", "gumroad", "payhip", "kofi", "pinterest"):
        assert (f / "listing" / f"guide-{m}.md").read_text().startswith("#")
    listing = json.loads((f / "listing" / "listing.json").read_text())
    assert len(listing["copy_"]["title"]) <= 140 and len(listing["copy_"]["tags"]) <= 13
    assert "AI-assisted" in listing["description"]
    with zipfile.ZipFile(f / "upload-package.zip") as z:
        names = z.namelist()
    assert any(n.startswith("product/") for n in names) and any(n.startswith("images/") for n in names)
    assert res.qa_status in ("pass", "warn") and not [c for c in res.checks if c.status == "fail"]
    assert "Done" in msgs[-1]


def test_catalogue_and_learning_loop(generated):
    res, _ = generated
    st = Store()
    assert st.product(res.product_id)["status"] == "draft"
    st.add_listing(res.product_id, "etsy", "https://etsy.com/x", 5.99)
    st.add_sale(res.product_id, "etsy", "2026-10-01", 2, 11.98)
    st.add_metric(res.product_id, "etsy", "2026-10-01", 120, 9, 3)
    row = st.performance()[0]
    assert (row["units"], round(row["revenue"], 2), row["views"], row["favourites"]) == (2, 11.98, 120, 9)
    assert st.product(res.product_id)["status"] == "listed"


def test_second_identical_product_is_flagged_not_blocked():
    a = generate(Request("teacher", "card_deck", sizes=["letter"], providers=["builtin"], seed=3, with_images=False))
    b = generate(Request("teacher", "card_deck", sizes=["letter"], providers=["builtin"], seed=3, with_images=False))
    assert b.similar[1] == a.product_id and b.similar[0] > 0.6
    assert any(c.name == "Uniqueness" and c.status == "warn" for c in b.checks)


def test_web_ui_pages_and_files(generated):
    res, _ = generated
    from moneyhub.web.app import app
    c = TestClient(app)
    for path in ["/", "/report", "/settings", f"/product/{res.product_id}"]:
        assert c.get(path).status_code == 200, path
    assert c.get(f"/files/{res.product_id}/images/etsy/01-hero.jpg").status_code == 200
    assert c.get(f"/files/{res.product_id}/../../catalog.sqlite").status_code == 404
    pdf = next((res.folder / "product").glob("*.pdf")).name
    assert c.get(f"/preview/{res.product_id}/product/{pdf}/0.png").status_code == 200
    r = c.post("/api/jobs", data={"niche": "adhd", "product_type": "card_deck", "llm": "builtin", "sizes": "letter"})
    assert r.status_code == 200 and len(r.json()["jobs"]) == 1


def test_spreadsheets_have_expected_structure():
    import warnings
    warnings.simplefilter("ignore")
    res = generate(Request("personal finance", "budget_tracker", sizes=["letter"], providers=["builtin"], seed=4, with_images=False))
    wb = openpyxl.load_workbook(next((res.folder / "product").glob("*.xlsx")))
    assert wb.sheetnames == ["Start Here", "Dashboard", "Settings", "Transactions", "Budget", "Savings", "Debt Payoff"]
    assert any(str(c.value).startswith("=SUMIFS") for row in wb["Budget"].iter_rows() for c in row)
    assert [c for c in res.checks if c.name.endswith("formulas")][0].status == "pass"


def test_niche_free_packs_get_volume_titles_and_only_match_on_same_design():
    def make(seed):
        return generate(Request("x", "clipart_pack", providers=["builtin"], palette="sage", seed=seed, with_images=False, with_video=False,
                                options={"no_research": True, "colourways": 1}))
    a, b, c = make(5), make(6), make(5)
    assert a.title == "Sage Boho Shapes Clipart Set" and "Vol." not in a.title
    assert "Vol. 2" in b.title and "Vol. 3" in c.title and b.title.endswith("Clipart Set")
    assert b.similar[0] == 0.0                        # same palette, different seed: a different design, no warning
    assert c.similar[0] == 1.0 and c.similar[1] == a.product_id   # same seed and options: a real duplicate
    assert any(x.name == "Uniqueness" and x.status == "warn" for x in c.checks)


def test_delete_removes_catalogue_row_and_folder_but_never_outside_the_output_dir(tmp_path):
    from moneyhub.pipeline import delete_product
    res = generate(Request("teacher", "card_deck", sizes=["letter"], providers=["builtin"], seed=3, with_images=False))
    assert res.folder.exists() and delete_product(res.product_id) and not res.folder.exists() and Store().product(res.product_id) is None
    keep = tmp_path / "precious"
    keep.mkdir()
    (keep / "a.txt").write_text("x")
    st = Store()
    pid = st.add_product(slug="x", title="t", niche="n", niche_name="n", product_type="card_deck", style="s", status="draft", folder=str(keep), price=1, qa=[],
                         fingerprint="", content_source="", listing_source="", spec="{}")
    assert delete_product(pid) and (keep / "a.txt").exists()   # the row goes, the foreign folder stays


def test_delete_returns_to_my_products_with_a_toast_message_and_uses_the_confirm_dialog(generated):
    res, _ = generated
    from moneyhub.web.app import app
    c = TestClient(app)
    page = c.get("/products").text
    assert f'action="/product/{res.product_id}/delete"' in page and "data-confirm-title" in page and "onsubmit" not in page   # no browser alert
    assert "confirm(" not in c.get(f"/product/{res.product_id}").text
    r = c.post(f"/product/{res.product_id}/delete", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/products?deleted=")
    assert Store().product(res.product_id) is None


def test_failed_generation_leaves_no_empty_folder(monkeypatch):
    from moneyhub.config import get_config
    from moneyhub.products.card_deck import CardDeck
    monkeypatch.setattr(CardDeck, "build", lambda self, ctx, content: (_ for _ in ()).throw(RuntimeError("boom")))
    with pytest.raises(RuntimeError):
        generate(Request("teacher", "card_deck", sizes=["letter"], providers=["builtin"], seed=3, with_images=False))
    products = get_config().output_dir / "products"
    assert not products.exists() or not any(products.iterdir())
